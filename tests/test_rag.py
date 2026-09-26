from unittest.mock import patch

import pytest
from fastapi.testclient import TestClient
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
import pymupdf

from app.api.main import app
from app.ingest.documents import parse_document
from app.ingest.chunker import chunk_document
from app.retrieval.answer import answer_question, NOT_FOUND
from app.retrieval.search import Hit
from app.providers.fake import FakeEmbeddingProvider


def test_office_locations(tmp_path):
    word = Document()
    word.add_heading('출장 규정', 1)
    word.add_paragraph('숙박비 상한은 100000원입니다.')
    path = tmp_path / 'rules.docx'
    word.save(path)
    chunks = chunk_document(parse_document(path))
    assert any('100000' in c.content and '문단' in c.section_path for c in chunks)
    assert all(c.page_from == 0 for c in chunks)
    book = Workbook()
    book.active.title = '예산'
    book.active.append(['출장', 100000])
    path = tmp_path / 'budget.xlsx'
    book.save(path)
    block = parse_document(path).blocks[0]
    assert block.section_path == '예산 > 행 1'
    assert ('B1', '100000') in block.pairs
    slides = Presentation()
    slide = slides.slides.add_slide(slides.slide_layouts[1])
    slide.shapes.title.text = '사내 규정'
    path = tmp_path / 'intro.pptx'
    slides.save(path)
    assert parse_document(path).blocks[0].page == 1


def test_table_unit_note_is_attached_to_every_row():
    """단위 표기는 표 밖 캡션에 있다. 행에 붙이지 않으면 1000배 틀린 답이 나온다."""
    from app.ingest.parser import ProseBlock, _table_blocks, unit_note

    prose = [ProseBlock('(단위:천원) 최근 3년간 발생한 매출 현황은 다음과 같습니다.', 5, '')]
    unit = unit_note(prose)
    assert unit == '단위:천원'

    class FakeTable:
        def extract(self):
            return [['구분', '2025년'], ['매출액', '9,876,543']]

    row = _table_blocks(FakeTable(), page_no=5, section='일반현황', unit=unit)[0]
    assert '9,876,543' in row.render() and '단위:천원' in row.render()
    # 단위 표기가 없는 페이지에서는 아무것도 붙이지 않는다.
    assert unit_note([ProseBlock('매출 현황', 5, '')]) == ''


def test_stream_does_not_warn_when_model_honestly_declines():
    """모델이 '찾지 못했다'고 답한 것은 정상 결과다. 인용 누락과 구분해야 한다."""
    from app.retrieval.answer import NOT_FOUND, answer_stream

    class Declining:
        name = 'declining'

        def generate_stream(self, system, user):
            yield '제공된 문서에서 찾지 못했습니다.'

    class Uncited:
        name = 'uncited'

        def generate_stream(self, system, user):
            yield '최대 8TB 입니다.'

    hit = Hit(1, 'Memory: 8TB', 'Memory', 2, 2, [], 'Guide', 'guide.pdf')
    done = {}
    for llm in (Declining(), Uncited()):
        with patch('app.retrieval.answer.search', return_value=[hit]):
            events = list(answer_stream(None, 'memory?', FakeEmbeddingProvider(), llm))
        done[llm.name] = [e for e in events if e['type'] == 'done'][0]

    assert done['declining']['answer'] == NOT_FOUND
    assert done['declining']['warning'] == ''          # 정직한 거절 -> 경고 없음
    assert done['uncited']['answer'] == NOT_FOUND
    assert '근거 번호' in done['uncited']['warning']    # 인용 누락 -> 경고


def test_vertically_merged_first_column_is_filled():
    """세로 병합 셀을 채우지 않으면 소계를 합계로 답한다.

    추출된 표에서 세로 병합은 첫 행에만 값이 있고 나머지는 비어 있다. 행 단위
    청크에서는 그 행이 무엇에 관한 값인지 잃는다. 실제로 "2025년 매출액"에
    시스템구축 소계(6,543,210)를 합계(9,876,543) 대신 답했다.
    """
    from app.ingest.parser import _table_blocks

    class FakeTable:
        def extract(self):
            return [
                ['구분', 'col1', '2025년'],
                ['매출액', '시스템구축', '6,543,210'],
                ['', '기술용역외', '3,333,333'],
                ['', '합계', '9,876,543'],
            ]

    rows = _table_blocks(FakeTable(), page_no=5, section='일반현황')
    rendered = [b.render() for b in rows]
    assert len(rendered) == 3
    assert all('매출액' in r for r in rendered), rendered
    총계 = [r for r in rendered if '합계' in r][0]
    assert '매출액' in 총계 and '9,876,543' in 총계


def test_pdf_page_and_long_text(tmp_path):
    path = tmp_path / 'guide.pdf'
    with pymupdf.open() as pdf:
        pdf.new_page().insert_text((50, 50), 'Memory maximum: 4TB.')
        pdf.save(path)
    assert chunk_document(parse_document(path))[0].page_from == 1
    from app.ingest.parser import ParsedDoc, ProseBlock
    doc = ParsedDoc('a', 'a', [], 1, 'hash', [ProseBlock('x' * 10000, 1, '')])
    chunks = chunk_document(doc)
    assert len(chunks) > 5
    assert all(len(c.content) <= 1202 for c in chunks)


@pytest.mark.parametrize('text,expected,count', [
    ('4TB [1]', '4TB [1]', 1),
    ('4TB [99]', NOT_FOUND, 0),
    ('4TB', NOT_FOUND, 0),
    (NOT_FOUND, NOT_FOUND, 0),
])
def test_citation_validation(text, expected, count):
    class LLM:
        name = 'test'
        def generate(self, system, user):
            return text
    hit = Hit(1, 'Memory: 4TB', 'Memory', 2, 2, [], 'Guide', 'guide.pdf')
    with patch('app.retrieval.answer.search', return_value=[hit]):
        result = answer_question(None, 'memory?', FakeEmbeddingProvider(), LLM())
    assert result.answer == expected
    assert len(result.sources) == count


def test_api_validation():
    client = TestClient(app)
    assert client.get('/').status_code == 200
    assert client.post('/ask', json={'question': '   '}).status_code == 422
    assert client.post('/ask', json={'question': 'hi', 'top_k': -1}).status_code == 422
    assert client.post('/ingest', files={'file': ('a.exe', b'bad')}).status_code == 415
    assert client.post('/ingest', files={'file': ('a.pdf', b'')}).status_code == 400


def test_cli_windows_glob(tmp_path):
    from app.ingest.cli import main
    for name in ['a.docx', 'b.docx']:
        d = Document()
        d.add_paragraph('Test document')
        d.save(tmp_path / name)
    assert main(['--dry-run', str(tmp_path / '*.docx')]) == 0
