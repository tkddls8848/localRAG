"""공통 문서 진입점. Office는 실제 위치를 섹션에 보존하며 페이지를 추정하지 않는다."""
from __future__ import annotations

import hashlib
from pathlib import Path

from app.ingest.parser import ParsedDoc, ProseBlock, TableRowBlock, _extract_products, parse_pdf

SUPPORTED = {'.pdf', '.docx', '.xlsx', '.pptx'}


def parse_document(path: str | Path) -> ParsedDoc:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED:
        raise ValueError(f'지원하지 않는 형식: {suffix}. PDF, DOCX, XLSX, PPTX를 사용하세요.')
    if suffix == '.pdf':
        return parse_pdf(path)
    doc = ParsedDoc(str(path), path.stem, _extract_products(path.stem), 0,
                    hashlib.sha256(path.read_bytes()).hexdigest())
    if suffix == '.docx':
        from docx import Document
        from docx.table import Table
        source = Document(path)
        doc.title = source.core_properties.title or path.stem
        section = '본문'
        for i, item in enumerate(source.iter_inner_content(), 1):
            if isinstance(item, Table):
                for r, row in enumerate(item.rows, 1):
                    doc.blocks.append(TableRowBlock(
                        [(f'열 {c}', cell.text) for c, cell in enumerate(row.cells, 1)],
                        0, f'{section} > 표 블록 {i}, 행 {r}'))
            elif item.text.strip():
                if item.style and item.style.name.startswith('Heading'):
                    section = item.text.strip()
                doc.blocks.append(ProseBlock(item.text, 0, f'{section} > 문단 {i}'))
    elif suffix == '.xlsx':
        from openpyxl import load_workbook
        source = load_workbook(path, read_only=True, data_only=False)
        try:
            for sheet in source:
                for r, row in enumerate(sheet.iter_rows(), 1):
                    pairs = [(cell.coordinate, str(cell.value)) for cell in row if cell.value is not None]
                    if pairs:
                        doc.blocks.append(TableRowBlock(pairs, 0, f'{sheet.title} > 행 {r}'))
        finally:
            source.close()
    else:
        from pptx import Presentation
        source = Presentation(path)
        doc.page_count = len(source.slides)
        for i, slide in enumerate(source.slides, 1):
            section = f'슬라이드 {i}'
            for shape in slide.shapes:
                if shape.has_text_frame and shape.text.strip():
                    doc.blocks.append(ProseBlock(shape.text, i, section))
                if shape.has_table:
                    for r, row in enumerate(shape.table.rows, 1):
                        doc.blocks.append(TableRowBlock(
                            [(f'열 {c}', cell.text) for c, cell in enumerate(row.cells, 1)],
                            i, f'{section} > 표 행 {r}'))
    doc.products = _extract_products(doc.title)
    return doc
